"""The markdown handed to the exporters is untrusted.

An article is written by a model from web search results, so anything in it — an image,
a stylesheet, a link — may have been put there by a stranger. Rendering it must never
make the worker reach the network or read its own disk, and must never embed a file in
the PDF. These tests aim the documented WeasyPrint vectors at a counting HTTP server and a
canary file, and assert that nothing was fetched and nothing was attached.
"""

from __future__ import annotations

import base64
import builtins
import http.server
import io
import threading
import zipfile

import pytest
from conftest import requires_pdf
from pypdf import PdfReader

from exporters import MissingDependencyError, to_docx, to_pdf

CANARY = "CANARY-7f3c-do-not-leak"

# A 1×1 transparent PNG — the one kind of resource an article may carry inline.
PNG_1PX = base64.b64encode(
    bytes.fromhex(
        "89504e470d0a1a0a0000000d4948445200000001000000010806000000"
        "1f15c4890000000d49444154789c6360000002000154a24f5d0000000049454e44ae426082"
    )
).decode()


class _CountingServer:
    """A local HTTP server that records every request it receives."""

    def __init__(self) -> None:
        hits: list[str] = []

        class Handler(http.server.BaseHTTPRequestHandler):
            def do_GET(self):
                hits.append(self.path)
                self.send_response(200)
                self.send_header("Content-Type", "text/plain")
                self.end_headers()
                self.wfile.write(CANARY.encode())

            def log_message(self, *args):  # keep test output quiet
                pass

        self.hits = hits
        self._httpd = http.server.ThreadingHTTPServer(("127.0.0.1", 0), Handler)
        self.url = f"http://127.0.0.1:{self._httpd.server_address[1]}"
        self._thread = threading.Thread(target=self._httpd.serve_forever, daemon=True)

    def __enter__(self):
        self._thread.start()
        return self

    def __exit__(self, *exc):
        self._httpd.shutdown()
        self._httpd.server_close()


@pytest.fixture()
def server():
    with _CountingServer() as running:
        yield running


@pytest.fixture()
def canary_file(tmp_path, monkeypatch):
    path = tmp_path / "canary.txt"
    path.write_text(CANARY)
    # Relative URLs used to resolve against the process's working directory.
    monkeypatch.chdir(tmp_path)
    return path


def _attachments(pdf_bytes: bytes) -> list[str]:
    """Every embedded file in the PDF, wherever it hangs.

    `PdfReader.attachments` only reads the document-level EmbeddedFiles tree, but
    `<a rel=attachment>` embeds through a link annotation instead — so this walks every
    object and decodes every stream, rather than trusting one entry point.
    """
    reader = PdfReader(io.BytesIO(pdf_bytes))
    found = list(reader.attachments)
    for number in range(1, int(reader.trailer["/Size"])):
        try:
            obj = reader.get_object(number)
        except Exception:  # free or unreadable xref entries
            continue
        if not hasattr(obj, "get"):
            continue
        if obj.get("/Type") == "/EmbeddedFile" or obj.get("/Subtype") == "/FileAttachment":
            found.append(f"object {number}")
        if hasattr(obj, "get_data"):
            try:
                if CANARY.encode() in obj.get_data():
                    found.append(f"canary in stream {number}")
            except Exception:
                continue
    return found


# Each payload is a way of making WeasyPrint fetch something. {url} is the counting server,
# {file} the canary on disk.
NETWORK_PAYLOADS = {
    "markdown image": "![x]({url}/md.png)",
    "raw img": '<img src="{url}/raw.png">',
    "img srcset": '<img srcset="{url}/set.png 1x">',
    "stylesheet link": '<link rel="stylesheet" href="{url}/s.css">',
    "css @import": '<style>@import url("{url}/import.css");</style>',
    "inline style url()": '<p style="background: url({url}/bg.png)">x</p>',
    "@font-face": "<style>@font-face {{ font-family: X; src: url({url}/f.ttf); }} p {{ font-family: X }}</style>\n\nx",
    "base href + relative img": '<base href="{url}/"><img src="relative.png">',
    "iframe, object, embed": '<iframe src="{url}/frame"></iframe><object data="{url}/obj"></object><embed src="{url}/e">',
    "inline svg image": '<svg width="10" height="10"><image href="{url}/svg.png" width="10" height="10"/></svg>',
    "svg data url with an external image": "![x](data:image/svg+xml;base64,{svg})",
    "attachment link to http": '<a rel="attachment" href="{url}/leak.txt">x</a>',
}

FILE_PAYLOADS = {
    "a rel=attachment file://": '<a rel="attachment" href="file://{file}">report</a>',
    "link rel=attachment file://": '<link rel="attachment" href="file://{file}">',
    "markdown attr_list attachment": '[report](file://{file}){{: rel="attachment"}}',
    "relative attachment": '<a rel="attachment" href="canary.txt">report</a>',
    "file:// image": '<img src="file://{file}">',
}


@requires_pdf
@pytest.mark.parametrize("payload", NETWORK_PAYLOADS.values(), ids=NETWORK_PAYLOADS.keys())
def test_rendering_never_reaches_the_network(server, payload):
    svg = base64.b64encode(
        f'<svg xmlns="http://www.w3.org/2000/svg"><image href="{server.url}/nested.png"/></svg>'.encode()
    ).decode()
    markdown = f"# Article\n\n{payload.format(url=server.url, svg=svg)}\n\nThe end."

    pdf_bytes = to_pdf(markdown)

    assert server.hits == [], f"the renderer fetched {server.hits}"
    assert _attachments(pdf_bytes) == []


@requires_pdf
@pytest.mark.parametrize("payload", FILE_PAYLOADS.values(), ids=FILE_PAYLOADS.keys())
def test_rendering_never_embeds_or_reads_a_local_file(canary_file, payload):
    markdown = f"# Article\n\n{payload.format(file=canary_file)}\n\nThe end."

    pdf_bytes = to_pdf(markdown)

    assert _attachments(pdf_bytes) == []
    assert CANARY.encode() not in pdf_bytes


@requires_pdf
def test_an_inline_png_still_renders():
    # The one resource an article may legitimately carry. Blocking everything else must not
    # take this with it.
    pdf_bytes = to_pdf(f"# Chart\n\n![chart](data:image/png;base64,{PNG_1PX})")
    images = [image for page in PdfReader(io.BytesIO(pdf_bytes)).pages for image in page.images]
    assert len(images) == 1


@requires_pdf
def test_links_stay_clickable():
    pdf_bytes = to_pdf("# Sources\n\nSee [the report](https://example.org/report).")
    annotations = [
        annotation.get_object()
        for page in PdfReader(io.BytesIO(pdf_bytes)).pages
        for annotation in page.get("/Annots", [])
    ]
    uris = [a["/A"]["/URI"] for a in annotations if "/A" in a and "/URI" in a["/A"]]
    assert uris == ["https://example.org/report"]


def test_a_missing_native_library_is_reported_not_raised_raw(monkeypatch):
    # Without Pango, `import weasyprint` raises OSError from cffi's dlopen — not ImportError.
    real_import = builtins.__import__

    def failing_import(name, *args, **kwargs):
        if name == "weasyprint" or name.startswith("weasyprint."):
            raise OSError("cannot load library 'libpango-1.0-0'")
        return real_import(name, *args, **kwargs)

    monkeypatch.setattr(builtins, "__import__", failing_import)
    with pytest.raises(MissingDependencyError, match="Pango"):
        to_pdf("# Title")


# ---------------------------------------------------------------- DOCX inline markup


def _docx_xml(docx_bytes: bytes) -> tuple[str, str]:
    with zipfile.ZipFile(io.BytesIO(docx_bytes)) as archive:
        return archive.read("word/document.xml").decode(), archive.read("word/_rels/document.xml.rels").decode()


def _paragraph_texts(docx_bytes: bytes) -> list[str]:
    import docx

    return [p.text for p in docx.Document(io.BytesIO(docx_bytes)).paragraphs]


def _italic_runs(docx_bytes: bytes) -> list[str]:
    import docx

    return [r.text for p in docx.Document(io.BytesIO(docx_bytes)).paragraphs for r in p.runs if r.italic]


def test_underscores_inside_a_word_are_not_emphasis():
    docx_bytes = to_docx("Call get_user_name with max_retry_count set.")
    assert _paragraph_texts(docx_bytes) == ["Call get_user_name with max_retry_count set."]
    assert _italic_runs(docx_bytes) == []


def test_spaced_asterisks_are_arithmetic_not_emphasis():
    docx_bytes = to_docx("So 2 * 3 = 6 and 4 * 5 = 20.")
    assert _paragraph_texts(docx_bytes) == ["So 2 * 3 = 6 and 4 * 5 = 20."]
    assert _italic_runs(docx_bytes) == []


def test_escaped_asterisks_stay_literal():
    assert _paragraph_texts(to_docx(r"A \*literal\* star.")) == ["A *literal* star."]


def test_real_emphasis_still_works():
    assert _italic_runs(to_docx("What comes *next* and _after_.")) == ["next", "after"]


def test_links_become_hyperlinks():
    document_xml, rels = _docx_xml(to_docx("Read [the report](https://example.org/report) today."))
    assert "<w:hyperlink" in document_xml
    assert 'Target="https://example.org/report"' in rels
    assert _paragraph_texts(to_docx("Read [the report](https://example.org/report) today.")) == [
        "Read the report today."
    ]


def test_unsafe_link_schemes_are_not_hyperlinked():
    docx_bytes = to_docx("Click [here](javascript:alert(1)) or [there](file:///etc/passwd).")
    document_xml, rels = _docx_xml(docx_bytes)
    assert "<w:hyperlink" not in document_xml
    assert "javascript:" not in rels and "file:" not in rels
    assert _paragraph_texts(docx_bytes) == ["Click here or there."]


def test_a_title_matching_the_first_heading_is_not_repeated():
    texts = _paragraph_texts(to_docx("# Quarterly Update\n\nBody.", title="Quarterly Update"))
    assert texts.count("Quarterly Update") == 1


def test_headings_are_inline_parsed():
    assert _paragraph_texts(to_docx("## The **APAC** plan")) == ["The APAC plan"]
