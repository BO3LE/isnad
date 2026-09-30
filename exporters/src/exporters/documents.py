"""markdown → PDF, markdown → DOCX.

Pure functions: same input, same output, nothing read or written outside the return
value. Optional dependencies (WeasyPrint, ``markdown``, python-docx) are imported
lazily inside each function so importing this module never requires them, and a
missing one raises `MissingDependencyError` instead of a bare `ImportError` from deep
inside a third-party package.
"""

from __future__ import annotations

import html as _html
import re
import zipfile
from io import BytesIO
from typing import TYPE_CHECKING

if TYPE_CHECKING:
    from docx.text.paragraph import Paragraph

_INLINE_PATTERN = re.compile(r"(\*\*.+?\*\*|`.+?`|\*.+?\*|_.+?_)")
_HEADING_RE = re.compile(r"^(#{1,6})\s+(.*)$")
_BULLET_RE = re.compile(r"^[-*+]\s+(.*)$")
_NUMBERED_RE = re.compile(r"^\d+\.\s+(.*)$")
_FENCE_RE = re.compile(r"^```")

_PDF_CSS = """
body { font-family: Georgia, 'Times New Roman', serif; color: #1a1a1a; line-height: 1.5; margin: 2.5cm; }
h1, h2, h3, h4, h5, h6 { font-family: Helvetica, Arial, sans-serif; color: #111111; }
code, pre { font-family: 'DejaVu Sans Mono', monospace; background: #f4f4f4; }
pre { padding: 0.75em; border-radius: 4px; overflow-x: auto; white-space: pre-wrap; }
"""

# zip entries store a modified-time on every part. python-docx's writer stamps that with
# the wall clock (see `_normalize_zip_timestamps`), which would otherwise make `to_docx`
# return different bytes for the same input depending on when it happens to run.
_FIXED_ZIP_TIMESTAMP = (1980, 1, 1, 0, 0, 0)


class MissingDependencyError(RuntimeError):
    """An optional rendering dependency isn't installed.

    Raised instead of letting a bare `ImportError` surface from inside WeasyPrint or
    python-docx, so callers get a message that says which extra to install.
    """


def to_pdf(markdown: str, *, title: str | None = None) -> bytes:
    """Render markdown to PDF bytes: `markdown` (the package) → HTML → WeasyPrint.

    Raises:
        TypeError: `markdown` is not a `str`.
        ValueError: `markdown` is empty, or WeasyPrint failed to lay it out.
        MissingDependencyError: the `pdf` extra (`markdown`, `weasyprint`) isn't installed.
    """
    _require_text(markdown)

    try:
        import markdown as markdown_lib
    except ImportError as exc:
        raise MissingDependencyError("to_pdf needs the 'markdown' package; install gp-exporters[pdf]") from exc
    try:
        from weasyprint import HTML
    except ImportError as exc:
        raise MissingDependencyError("to_pdf needs WeasyPrint; install gp-exporters[pdf]") from exc

    body_html = markdown_lib.markdown(markdown, extensions=["extra", "sane_lists"])
    title_tag = f"<title>{_html.escape(title)}</title>" if title else ""
    document_html = (
        "<!DOCTYPE html><html><head><meta charset='utf-8'>"
        f"{title_tag}<style>{_PDF_CSS}</style></head><body>{body_html}</body></html>"
    )

    try:
        pdf_bytes = HTML(string=document_html, base_url=".").write_pdf()
    except Exception as exc:  # WeasyPrint's own errors vary by failure; wrap them uniformly
        raise ValueError(f"markdown could not be rendered to PDF: {exc}") from exc
    if not pdf_bytes:
        raise ValueError("WeasyPrint returned no PDF bytes")
    return pdf_bytes


def to_docx(markdown: str, *, title: str | None = None) -> bytes:
    """Render markdown to DOCX bytes with python-docx.

    Handles headings (`#` … `######`), bold (`**text**`), italic (`*text*` / `_text_`),
    inline code (`` `code` ``), bullet and numbered lists, and fenced code blocks
    (```` ``` ````). Anything else is emitted as a plain paragraph.

    Raises:
        TypeError: `markdown` is not a `str`.
        ValueError: `markdown` is empty.
        MissingDependencyError: the `docx` extra (python-docx) isn't installed.
    """
    _require_text(markdown)

    try:
        import docx
        from docx.shared import Inches
    except ImportError as exc:
        raise MissingDependencyError("to_docx needs python-docx; install gp-exporters[docx]") from exc

    document = docx.Document()
    if title:
        document.core_properties.title = title
        document.add_heading(title, level=0)

    code_lines: list[str] = []
    in_code_block = False

    def flush_code_block() -> None:
        if not code_lines:
            return
        code_paragraph = document.add_paragraph()
        code_paragraph.paragraph_format.left_indent = Inches(0.3)
        code_run = code_paragraph.add_run("\n".join(code_lines))
        code_run.font.name = "Consolas"
        code_lines.clear()

    for raw_line in markdown.splitlines():
        line = raw_line.rstrip()

        if _FENCE_RE.match(line):
            if in_code_block:
                flush_code_block()
            in_code_block = not in_code_block
            continue
        if in_code_block:
            code_lines.append(raw_line)
            continue
        if not line.strip():
            continue

        heading_match = _HEADING_RE.match(line)
        if heading_match:
            level = len(heading_match.group(1))
            document.add_heading(heading_match.group(2).strip(), level=level)
            continue

        bullet_match = _BULLET_RE.match(line)
        if bullet_match:
            _add_inline_runs(document.add_paragraph(style="List Bullet"), bullet_match.group(1))
            continue

        numbered_match = _NUMBERED_RE.match(line)
        if numbered_match:
            _add_inline_runs(document.add_paragraph(style="List Number"), numbered_match.group(1))
            continue

        _add_inline_runs(document.add_paragraph(), line)

    # An unterminated ``` fence still renders whatever it captured, rather than
    # silently dropping it — malformed input degrades gracefully.
    if in_code_block:
        flush_code_block()

    buffer = BytesIO()
    document.save(buffer)
    return _normalize_zip_timestamps(buffer.getvalue())


def _require_text(markdown: str) -> None:
    if not isinstance(markdown, str):
        raise TypeError(f"markdown must be a str, got {type(markdown).__name__}")
    if not markdown.strip():
        raise ValueError("markdown must not be empty")


def _add_inline_runs(paragraph: Paragraph, text: str) -> None:
    """Split `text` on **bold**, *italic*/_italic_ and `code` spans and add matching runs."""
    for token in _INLINE_PATTERN.split(text):
        if not token:
            continue
        if token.startswith("**") and token.endswith("**") and len(token) > 4:
            paragraph.add_run(token[2:-2]).bold = True
        elif token.startswith("`") and token.endswith("`") and len(token) > 2:
            paragraph.add_run(token[1:-1]).font.name = "Consolas"
        elif len(token) > 2 and (
            (token.startswith("*") and token.endswith("*")) or (token.startswith("_") and token.endswith("_"))
        ):
            paragraph.add_run(token[1:-1]).italic = True
        else:
            paragraph.add_run(token)


def _normalize_zip_timestamps(data: bytes) -> bytes:
    """Rewrite every zip entry's embedded timestamp to a fixed value.

    python-docx's OPC writer stamps each part with `time.localtime()` (the default
    `zipfile.ZipInfo.date_time` when none is given to `writestr`), which would otherwise
    make `to_docx` return different bytes for the same input depending on the moment it
    runs — breaking both the "pure function" rule and byte-exact golden-file testing.
    """
    source = zipfile.ZipFile(BytesIO(data))
    out_buffer = BytesIO()
    with zipfile.ZipFile(out_buffer, "w", zipfile.ZIP_DEFLATED) as out:
        for info in source.infolist():
            normalized = zipfile.ZipInfo(info.filename, date_time=_FIXED_ZIP_TIMESTAMP)
            normalized.compress_type = info.compress_type
            normalized.external_attr = info.external_attr
            out.writestr(normalized, source.read(info.filename))
    return out_buffer.getvalue()
